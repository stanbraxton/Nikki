"""Document tools: read any office/PDF/image file in the workspace, create and edit Excel
workbooks, build PowerPoint decks, create PDFs from Markdown, fill PDF forms and stamp
signatures. Every created file is announced with the artifacts FILE_MARK so the chat
attaches it for the user.

All paths are relative to Nikki's workspace (uploads land in `uploads/`).
"""
from __future__ import annotations

import base64
import csv
import io
import json
import logging
import time
from pathlib import Path

from langchain_core.tools import tool

from app.config import settings
from app.tools.artifacts import out_dir, safe_name, saved, workspace_path

log = logging.getLogger("nikki.documents")

MAX_CHARS = 60_000
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp"}
AUDIO_EXT = {".mp3", ".m4a", ".wav", ".webm", ".ogg", ".oga", ".mp4", ".mpeg", ".mpga", ".flac"}


def _clip(text: str, limit: int = MAX_CHARS) -> str:
    if len(text) > limit:
        return text[:limit] + f"\n\n[... truncated, {len(text) - limit} more chars]"
    return text


# ---------------------------------------------------------------- readers
def read_pdf(p: Path, max_pages: int = 150) -> str:
    import pypdf

    reader = pypdf.PdfReader(str(p))
    parts = []
    for i, page in enumerate(reader.pages[:max_pages], 1):
        t = page.extract_text() or ""
        if t.strip():
            parts.append(f"--- Page {i} ---\n{t}")
    fields = reader.get_fields() or {}
    head = f"[PDF, {len(reader.pages)} pages{', fillable form with ' + str(len(fields)) + ' fields' if fields else ''}]\n"
    return head + ("\n\n".join(parts) if parts else "(no extractable text — scanned image? try describing it via an image export)")


def read_docx(p: Path) -> str:
    import docx

    d = docx.Document(str(p))
    out = []
    for para in d.paragraphs:
        if para.text.strip():
            style = (para.style.name or "") if para.style else ""
            prefix = "# " if style.startswith("Heading 1") or style == "Title" else "## " if style.startswith("Heading") else ""
            out.append(prefix + para.text)
    for ti, table in enumerate(d.tables, 1):
        out.append(f"\n[Table {ti}]")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out) or "(empty document)"


def read_xlsx(p: Path, max_rows: int = 400) -> str:
    import openpyxl

    wb = openpyxl.load_workbook(str(p), data_only=True)
    wbf = openpyxl.load_workbook(str(p))  # formulas, for cells whose cached value is missing
    out = []
    for ws in wb.worksheets:
        wsf = wbf[ws.title]
        out.append(f"=== Sheet: {ws.title} ({ws.max_row} rows x {ws.max_column} cols) ===")
        n = 0
        for row in ws.iter_rows():
            if all(c.value is None for c in row):
                continue
            vals = []
            for c in row:
                v = c.value
                if v is None:
                    f = wsf[c.coordinate].value
                    v = f if isinstance(f, str) and f.startswith("=") else None
                vals.append("" if v is None else str(v))
            out.append(" | ".join(vals))
            n += 1
            if n >= max_rows:
                out.append(f"... (more rows not shown; first {max_rows} non-empty rows listed)")
                break
    return "\n".join(out)


def read_pptx(p: Path) -> str:
    from pptx import Presentation

    prs = Presentation(str(p))
    out = []
    for i, slide in enumerate(prs.slides, 1):
        out.append(f"--- Slide {i} ---")
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                out.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    out.append(" | ".join(c.text for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out.append(f"[notes] {slide.notes_slide.notes_text_frame.text}")
    return "\n".join(out) or "(no slides)"


def describe_image(p: Path, question: str = "") -> str:
    """Vision read of an image via Claude (Anthropic). Transcribes text exactly, then describes."""
    if not settings.anthropic_api_key:
        return "error: ANTHROPIC_API_KEY is not configured, cannot look at images"
    from anthropic import Anthropic

    mime = {"png": "image/png", "jpg": "image/jpeg", "jpeg": "image/jpeg", "gif": "image/gif", "webp": "image/webp"}.get(p.suffix.lower().lstrip("."), "image/png")
    b64 = base64.b64encode(p.read_bytes()).decode()
    ask = question.strip() or (
        "First transcribe ALL visible text, numbers, labels and table contents exactly as written, preserving layout "
        "(use a Markdown table for tabular data). Then describe what the image shows in 2-4 sentences."
    )
    client = Anthropic(api_key=settings.anthropic_api_key)
    r = client.messages.create(
        model="claude-3-5-sonnet-20241022",
        max_tokens=4096,
        messages=[{"role": "user", "content": [
            {"type": "image", "source": {"type": "base64", "media_type": mime, "data": b64}},
            {"type": "text", "text": ask}
        ]}],
    )
    return r.content[0].text if r.content else "(no description)"


def read_any(p: Path, question: str = "") -> str:
    """Dispatch by extension. Used by read_document and by the chat upload handler."""
    ext = p.suffix.lower()
    if ext == ".pdf":
        return read_pdf(p)
    if ext in (".docx", ".docm"):
        return read_docx(p)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(p)
    if ext == ".pptx":
        return read_pptx(p)
    if ext in IMAGE_EXT:
        return describe_image(p, question)
    if ext in AUDIO_EXT:
        from app.tools.speech import transcribe_path

        return transcribe_path(p)
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(p.read_text(encoding="utf-8", errors="replace"))))
        return "\n".join(" | ".join(r) for r in rows[:500]) + (f"\n... ({len(rows)} rows total)" if len(rows) > 500 else "")
    if ext in (".txt", ".md", ".json", ".xml", ".yaml", ".yml", ".html", ".htm", ".py", ".js", ".ts", ".log", ".ini", ".toml"):
        return p.read_text(encoding="utf-8", errors="replace")
    if ext == ".doc":
        return "error: legacy .doc is not supported — ask for a .docx or PDF export"
    return f"error: unsupported file type {ext}"


@tool
def read_document(path: str, question: str = "") -> str:
    """Read a file in the workspace and return its contents as text: PDF (text + form-field count),
    Word .docx (headings, paragraphs, tables), Excel .xlsx (every sheet, row by row), PowerPoint .pptx
    (per-slide text + notes), CSV/TXT/MD/JSON, images (text transcribed exactly + description; pass
    `question` to ask something specific about the picture) and audio (transcribed).
    Files the user attaches in chat are saved under uploads/ — their paths are given in the message."""
    try:
        p = workspace_path(path, must_exist=True)
        return _clip(read_any(p, question))
    except Exception as e:  # noqa: BLE001 — return errors to the model, never raise
        return f"error reading {path}: {type(e).__name__}: {e}"


# ---------------------------------------------------------------- Excel
def _parse_json(s: str, what: str):
    try:
        return json.loads(s)
    except json.JSONDecodeError as e:
        raise ValueError(f"{what} must be valid JSON: {e}") from e


@tool
def xlsx_create(name: str, sheets_json: str) -> str:
    """Create an Excel workbook in the workspace (documents/ folder). `sheets_json` is a JSON object
    mapping sheet name -> list of rows (each row a list of cells; numbers stay numeric, strings starting
    with '=' become formulas). The first row is styled as a bold header and columns are auto-sized.
    Example: {"Budget": [["Item","Cost"],["Rent",1200],["Total","=SUM(B2:B2)"]]}"""
    import openpyxl
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    try:
        sheets = _parse_json(sheets_json, "sheets_json")
        if not isinstance(sheets, dict) or not sheets:
            return "error: sheets_json must be a non-empty object of sheetName -> rows"
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for title, rows in sheets.items():
            ws = wb.create_sheet(title=str(title)[:31])
            widths: dict[int, int] = {}
            for r in rows:
                ws.append(list(r))
                for i, v in enumerate(r, 1):
                    widths[i] = max(widths.get(i, 8), min(len(str(v)) + 2, 60))
            for c in ws[1] if rows else []:
                c.font = Font(bold=True)
            for i, w in widths.items():
                ws.column_dimensions[get_column_letter(i)].width = w
            ws.freeze_panes = "A2"
        fname = safe_name(name, "workbook")
        if not fname.lower().endswith(".xlsx"):
            fname += ".xlsx"
        p = out_dir("documents") / fname
        wb.save(str(p))
        return saved(p, f"{len(sheets)} sheet(s)")
    except Exception as e:  # noqa: BLE001
        return f"error: {type(e).__name__}: {e}"


@tool
def xlsx_update(path: str, cells_json: str, sheet: str = "") -> str:
    """Edit an existing Excel workbook in place. `cells_json` is a JSON object of cell address -> value
    (e.g. {"B2": 1500, "C2": "=B2*1.08", "A10": "Note"}); `sheet` defaults to the first sheet.
    Use read_document first to see the current layout. Formulas recalculate when the file is opened."""
    import openpyxl

    try:
        p = workspace_path(path, must_exist=True)
        cells = _parse_json(cells_json, "cells_json")
        if not isinstance(cells, dict):
            return "error: cells_json must be an object of cellAddress -> value"
        wb = openpyxl.load_workbook(str(p))
        ws = wb[sheet] if sheet else wb.worksheets[0]
        for addr, val in cells.items():
            ws[addr] = val
        wb.save(str(p))
        return saved(p, f"updated {len(cells)} cell(s) on '{ws.title}'")
    except Exception as e:  # noqa: BLE001
        return f"error: {type(e).__name__}: {e}"


# ---------------------------------------------------------------- PowerPoint
@tool
def pptx_create(name: str, slides_json: str) -> str:
    """Create a PowerPoint deck in the workspace (documents/ folder). `slides_json` is a JSON list of
    slides: {"title": "...", "bullets": ["...", "..."], "notes": "optional speaker notes",
    "image": "optional workspace path to a picture"}. The first slide is rendered as a title slide
    when it has no bullets (its first note line becomes the subtitle). 16:9, clean default theme."""
    from pptx import Presentation
    from pptx.util import Inches, Pt

    try:
        slides = _parse_json(slides_json, "slides_json")
        if not isinstance(slides, list) or not slides:
            return "error: slides_json must be a non-empty list"
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        for i, s in enumerate(slides):
            title = str(s.get("title", ""))
            bullets = [str(b) for b in (s.get("bullets") or [])]
            if i == 0 and not bullets:
                sl = prs.slides.add_slide(prs.slide_layouts[0])
                sl.shapes.title.text = title
                if len(sl.placeholders) > 1:
                    sl.placeholders[1].text = str(s.get("subtitle") or (s.get("notes") or "").split("\n")[0])
            else:
                sl = prs.slides.add_slide(prs.slide_layouts[1])
                sl.shapes.title.text = title
                body = sl.placeholders[1].text_frame
                body.clear()
                for j, b in enumerate(bullets):
                    para = body.paragraphs[0] if j == 0 else body.add_paragraph()
                    level = 0
                    while b.startswith("  "):
                        b, level = b[2:], level + 1
                    para.text = b.lstrip("-• ")
                    para.level = min(level, 4)
                    para.font.size = Pt(24 if level == 0 else 20)
                if s.get("image"):
                    img = workspace_path(str(s["image"]), must_exist=True)
                    sl.placeholders[1].width = Inches(6.5)
                    sl.shapes.add_picture(str(img), Inches(7.3), Inches(1.8), width=Inches(5.5))
            if s.get("notes"):
                sl.notes_slide.notes_text_frame.text = str(s["notes"])
        fname = safe_name(name, "deck")
        if not fname.lower().endswith(".pptx"):
            fname += ".pptx"
        p = out_dir("documents") / fname
        prs.save(str(p))
        return saved(p, f"{len(slides)} slide(s)")
    except Exception as e:  # noqa: BLE001
        return f"error: {type(e).__name__}: {e}"


# ---------------------------------------------------------------- PDF
def _pdf_class():
    from fpdf import FPDF

    class Doc(FPDF):
        def footer(self):
            self.set_y(-15)
            self.set_font("Helvetica", "I", 8)
            self.set_text_color(120)
            self.cell(0, 10, f"Page {self.page_no()}", align="C")

    return Doc


def _md_to_html(md: str) -> str:
    """Tiny Markdown subset -> HTML fpdf2 understands (headings, bullets, bold/italic, tables, hr)."""
    import html as _h
    import re

    lines = md.splitlines()
    out, in_ul, in_ol, table = [], False, False, []

    def inline(s: str) -> str:
        s = _h.escape(s)
        s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
        s = re.sub(r"(?<!\*)\*(?!\*)(.+?)\*", r"<i>\1</i>", s)
        s = re.sub(r"`(.+?)`", r"<font face='Courier'>\1</font>", s)
        s = re.sub(r"\[(.+?)\]\((https?://[^)]+)\)", r"<a href='\2'>\1</a>", s)
        return s

    def flush_lists():
        nonlocal in_ul, in_ol
        if in_ul:
            out.append("</ul>")
            in_ul = False
        if in_ol:
            out.append("</ol>")
            in_ol = False

    def flush_table():
        nonlocal table
        if not table:
            return
        rows = [r for r in table if not re.match(r"^\s*\|?\s*:?-{2,}", r)]
        cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        if cells:
            ncol = max(len(r) for r in cells)
            w = int(100 / ncol)
            out.append("<table border='1' width='100%'><thead><tr>" + "".join(f"<th width='{w}%'>{inline(c)}</th>" for c in cells[0]) + "</tr></thead><tbody>")
            for r in cells[1:]:
                out.append("<tr>" + "".join(f"<td>{inline(c)}</td>" for c in r + [""] * (ncol - len(r))) + "</tr>")
            out.append("</tbody></table><br>")
        table = []

    for ln in lines:
        if ln.strip().startswith("|"):
            flush_lists()
            table.append(ln)
            continue
        flush_table()
        m = re.match(r"^(#{1,4})\s+(.*)", ln)
        if m:
            flush_lists()
            lvl = len(m.group(1))
            out.append(f"<h{lvl}>{inline(m.group(2))}</h{lvl}>")
        elif re.match(r"^\s*[-*•]\s+", ln):
            if in_ol:
                out.append("</ol>")
                in_ol = False
            if not in_ul:
                out.append("<ul>")
                in_ul = True
            out.append(f"<li>{inline(re.sub(r'^\s*[-*•]\s+', '', ln))}</li>")
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            if in_ul:
                out.append("</ul>")
                in_ul = False
            if not in_ol:
                out.append("<ol>")
                in_ol = True
            out.append(f"<li>{inline(re.sub(r'^\s*\d+[.)]\s+', '', ln))}</li>")
        elif re.match(r"^\s*(-{3,}|\*{3,})\s*$", ln):
            flush_lists()
            out.append("<hr>")
        elif not ln.strip():
            flush_lists()
            out.append("<br>")
        else:
            flush_lists()
            out.append(f"<p>{inline(ln)}</p>")
    flush_lists()
    flush_table()
    return "\n".join(out)


@tool
def pdf_create(name: str, markdown: str, title: str = "") -> str:
    """Create a PDF document in the workspace (documents/ folder) from Markdown text: headings (#, ##),
    paragraphs, bullet and numbered lists, **bold**, *italic*, links, pipe tables and --- rules.
    `title` (optional) is printed as a large heading on the first page. Use for reports, letters,
    outlines and anything the user wants as a PDF."""
    try:
        Doc = _pdf_class()
        pdf = Doc(orientation="P", unit="mm", format="Letter")
        pdf.set_margins(20, 20, 20)
        pdf.set_auto_page_break(auto=True, margin=20)
        pdf.add_page()
        pdf.set_font("Helvetica", size=11)
        html = ""
        if title:
            html += f"<h1>{title}</h1>"
        html += _md_to_html(markdown)
        # fpdf2 core fonts are latin-1: replace unsupported glyphs rather than crash
        html = html.replace("\u2014", "-").replace("\u2013", "-").replace("\u2019", "'").replace("\u201c", '"').replace("\u201d", '"').replace("\u2026", "...")
        html = html.encode("latin-1", "replace").decode("latin-1")
        pdf.write_html(html)
        fname = safe_name(name, "document")
        if not fname.lower().endswith(".pdf"):
            fname += ".pdf"
        p = out_dir("documents") / fname
        pdf.output(str(p))
        return saved(p, f"{pdf.page_no()} page(s)")
    except Exception as e:  # noqa: BLE001
        return f"error: {type(e).__name__}: {e}"


@tool
def pdf_form_fields(path: str) -> str:
    """List the fillable form fields of a PDF (name, type, current value, checkbox options) so you
    can fill them with pdf_fill_form. Returns 'no fillable fields' for flat PDFs."""
    import pypdf

    try:
        p = workspace_path(path, must_exist=True)
        reader = pypdf.PdfReader(str(p))
        fields = reader.get_fields() or {}
        if not fields:
            return "no fillable fields in this PDF (it is a flat document; pdf_sign can still stamp text on it)"
        rows = []
        for name, f in fields.items():
            ft = str(f.get("/FT", "")).strip("/")
            kind = {"Tx": "text", "Btn": "checkbox/radio", "Ch": "choice", "Sig": "signature"}.get(ft, ft or "?")
            val = f.get("/V", "")
            states = ""
            if ft == "Btn":
                ap = f.get("/_States_") or []
                if ap:
                    states = f" options={list(ap)}"
            rows.append(f"- {name} [{kind}]{states} value={val!r}")
        return f"{len(fields)} fields:\n" + "\n".join(rows)
    except Exception as e:  # noqa: BLE001
        return f"error: {type(e).__name__}: {e}"


@tool
def pdf_fill_form(path: str, values_json: str, output_name: str = "") -> str:
    """Fill a PDF form. `values_json` maps field name -> value (text for text fields; for checkboxes
    use the option name from pdf_form_fields, e.g. "/Yes", or true/false). Writes a new file in
    documents/ (default '<original>-filled.pdf') and leaves the original untouched."""
    import pypdf

    try:
        p = workspace_path(path, must_exist=True)
        values = _parse_json(values_json, "values_json")
        if not isinstance(values, dict):
            return "error: values_json must be an object of fieldName -> value"
        reader = pypdf.PdfReader(str(p))
        fields = reader.get_fields() or {}
        unknown = [k for k in values if k not in fields]
        writer = pypdf.PdfWriter()
        writer.append(reader)
        norm: dict[str, object] = {}
        for k, v in values.items():
            if k not in fields:
                continue
            f = fields[k]
            if str(f.get("/FT", "")) == "/Btn":
                states = [s for s in (f.get("/_States_") or []) if s != "/Off"]
                if v is True or (isinstance(v, str) and v.lower() in ("true", "yes", "x", "on")):
                    norm[k] = states[0] if states else "/Yes"
                elif v is False or (isinstance(v, str) and v.lower() in ("false", "no", "off", "")):
                    norm[k] = "/Off"
                else:
                    norm[k] = v if str(v).startswith("/") else f"/{v}"
            else:
                norm[k] = "" if v is None else str(v)
        for page in writer.pages:
            writer.update_page_form_field_values(page, norm, auto_regenerate=False)
        writer.set_need_appearances_writer(True)
        fname = safe_name(output_name) if output_name else f"{p.stem}-filled"
        if not fname.lower().endswith(".pdf"):
            fname += ".pdf"
        out = out_dir("documents") / fname
        with out.open("wb") as fh:
            writer.write(fh)
        note = f"filled {len(norm)} field(s)"
        if unknown:
            note += f"; ignored unknown fields: {unknown}"
        return saved(out, note)
    except Exception as e:  # noqa: BLE001
        return f"error: {type(e).__name__}: {e}"


@tool
def pdf_sign(path: str, signer_name: str, page: int = 1, x_mm: float = 25, y_mm: float = 250,
             width_mm: float = 60, signature_image: str = "", date_line: bool = True, output_name: str = "") -> str:
    """Stamp a visible signature on a PDF page: either a signature image from the workspace
    (`signature_image`, PNG with transparent background works best) or a script-style rendering of
    `signer_name`; a printed name and today's date go underneath (`date_line`). Position is from the
    top-left of the page in millimetres (Letter page = 216 x 279 mm; y_mm=250 is near the bottom).
    Writes documents/<original>-signed.pdf. This is a visual signature, not a cryptographic one."""
    import pypdf
    from fpdf import FPDF

    try:
        p = workspace_path(path, must_exist=True)
        reader = pypdf.PdfReader(str(p))
        if not 1 <= page <= len(reader.pages):
            return f"error: page must be 1..{len(reader.pages)}"
        target = reader.pages[page - 1]
        w_pt, h_pt = float(target.mediabox.width), float(target.mediabox.height)
        overlay = FPDF(unit="pt", format=(w_pt, h_pt))
        overlay.add_page()
        x, y, w = x_mm * 72 / 25.4, y_mm * 72 / 25.4, width_mm * 72 / 25.4
        if signature_image:
            img = workspace_path(signature_image, must_exist=True)
            overlay.image(str(img), x=x, y=y, w=w)
            base_y = y + w * 0.35
        else:
            overlay.set_font("Times", "I", 26)
            overlay.set_text_color(20, 30, 90)
            overlay.text(x, y + 22, signer_name)
            base_y = y + 30
        overlay.set_draw_color(60)
        overlay.line(x, base_y, x + w, base_y)
        overlay.set_font("Helvetica", size=9)
        overlay.set_text_color(40)
        line = signer_name + (f"    Date: {time.strftime('%m/%d/%Y')}" if date_line else "")
        overlay.text(x, base_y + 12, line)
        ov_reader = pypdf.PdfReader(io.BytesIO(overlay.output()))
        writer = pypdf.PdfWriter()
        writer.append(reader)
        writer.pages[page - 1].merge_page(ov_reader.pages[0])
        fname = safe_name(output_name) if output_name else f"{p.stem}-signed"
        if not fname.lower().endswith(".pdf"):
            fname += ".pdf"
        out = out_dir("documents") / fname
        with out.open("wb") as fh:
            writer.write(fh)
        return saved(out, f"signed page {page}")
    except Exception as e:  # noqa: BLE001
        return f"error: {type(e).__name__}: {e}"


for _t in (read_document, pdf_form_fields, xlsx_create, pptx_create, pdf_create, pdf_fill_form, pdf_sign):
    _t.metadata = {"requires_approval": False}  # read-only, or they only create NEW files inside the workspace
xlsx_update.metadata = {"requires_approval": True}  # edits an existing file in place

TOOLS = [read_document, xlsx_create, xlsx_update, pptx_create, pdf_create, pdf_form_fields, pdf_fill_form, pdf_sign]
