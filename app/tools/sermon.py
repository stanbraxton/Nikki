"""Sermon-outline compiler: renders Pastor Stan's exact Sermon Outline Template as a .docx and drops it in
Google Drive Church/Sermons. The layout is fixed in code (same renderer Viktor used for "It Ends at God" and
"Worth the Dig"); the model only supplies content. See knowledge/sermon-prep.md for the content rules."""
from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

from langchain_core.tools import tool

from app.config import settings

RULE = "──────────────────────────────"
SERMONS_FOLDER_PATH = "Church/Sermons"

SCHEMA = {
    "series": "Kingdom of God (Movement Two: Knowledge)",
    "title": "Worth the Dig",
    "title_logic": "✍️ why this title is a promise + hook, plus 2-3 alternates (optional)",
    "passage_ref": "Proverbs 2:1–6 (NKJV)",
    "passage_text": "full NKJV text of the passage",
    "supporting_texts": "Hosea 4:6; Hebrews 11:6; ...",
    "goal": "one-sentence message goal",
    "intro": {
        "hook_title": "The Hook — <name of the story>",
        "hook": ["paragraph", "paragraph", "..."],
        "plant_list": "(Plant list — details said cold in the intro that the points harvest; what the open loop is)",
        "confusions": ["1. <Area> (psychological/social): ...", "2. <Area> (intellectual): ...", "3. <Area> (church-wide): ..."],
    },
    "structure_note": "e.g. IF (vv1–4) → THEN (v5) → FOR (v6) — optional",
    "points": [
        {
            "heading": "1. MAIN POINT IN ALL CAPS (verses)",
            "statement": "short declarative sentence linking point to text",
            "bullets": ["Sub-Concept A (...): ...", "Sub-Concept B (...): ..."],
            "sub_bullets": {"0": ["optional level-2 bullets under bullet index 0", "..."]},
            "note": "optional italic note after the point",
        }
    ],
    "illustrations": [{"heading": "ILLUSTRATION 1: <title> (explains Point 1)", "paragraphs": ["..."]}],
    "background": ["paragraph", "paragraph"],
    "conclusion": {
        "recap": "summary / final recap",
        "gospel_connection": "...",
        "invitation": "...",
        "today": "Immediate Response (Today): ...",
        "this_week": "Weekly Integration (This Week): ...",
    },
    "takeaways": ["line", "line"],
    "soundbites": ['"quote"', '"quote"'],
    "preaching_notes": ["optional coaching reminders accepted in the session"],
}


def _slug(s: str) -> str:
    s = re.sub(r"[^\w\s-]", "", s).strip()
    return re.sub(r"\s+", " ", s) or "Sermon"


def render_docx(o: dict[str, Any], out: Path) -> Path:
    from docx import Document
    from docx.shared import Pt

    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    def P(t: str, style: str | None = None, bold: bool = False, italic: bool = False):
        p = d.add_paragraph(style=style) if style else d.add_paragraph()
        r = p.add_run(t)
        r.bold, r.italic = bold, italic
        return p

    def H1(t: str):
        d.add_heading(t, 1)

    def H2(t: str):
        d.add_heading(t, 2)

    def B(t: str, lvl: int = 1):
        P(t, "List Bullet" if lvl == 1 else "List Bullet 2")

    def rule():
        P(RULE)

    def paras(x: Any, **kw):
        if isinstance(x, str):
            x = [x]
        for t in x or []:
            if t:
                P(t, **kw)

    title = (o.get("title") or "").replace("✍️", "").strip()
    d.add_heading(f"SERMON OUTLINE — {title.upper()}", 0)
    P("Note: Everything unmarked came from Pastor Stan's own study and words in the prep session. Lines marked ✍️ were "
      "drafted by Nikki from that material to fill the template — rework them into your voice.", italic=True)

    H2(f"SERIES / MINISTRY BRANDING: {o.get('series', '')}")
    H2(f"MESSAGE TITLE: {o.get('title', '')}")
    if o.get("title_logic"):
        P(o["title_logic"], italic=True)
    H2(f"SUGGESTED PASSAGE: {o.get('passage_ref', '')}")
    paras(o.get("passage_text"))
    if o.get("supporting_texts"):
        P(f"Supporting texts: {o['supporting_texts']}")
    P("MESSAGE GOAL:", bold=True)
    paras(o.get("goal"))
    rule()

    intro = o.get("intro") or {}
    H1("INTRODUCTION")
    if intro.get("hook_title"):
        P(intro["hook_title"], bold=True)
    paras(intro.get("hook"))
    if intro.get("plant_list"):
        P(intro["plant_list"], italic=True)
    P("The cultural confusion or practical tension regarding this issue can be seen across these specific areas in society today:")
    paras(intro.get("confusions"))
    rule()

    H1("SERMON POINTS")
    if o.get("structure_note"):
        P(o["structure_note"], italic=True)
    for pt in o.get("points") or []:
        H2(pt.get("heading", ""))
        paras(pt.get("statement"))
        subs = pt.get("sub_bullets") or {}
        for i, b in enumerate(pt.get("bullets") or []):
            B(b)
            for sb in subs.get(str(i), subs.get(i, [])) or []:
                B(sb, 2)
        if pt.get("note"):
            P(pt["note"], italic=True)
    rule()

    H1("SERMON ILLUSTRATIONS")
    for il in o.get("illustrations") or []:
        H2(il.get("heading", ""))
        paras(il.get("paragraphs"))
    rule()

    H1("BACKGROUND BIBLICAL HISTORY AND CULTURE")
    paras(o.get("background"))
    rule()

    c = o.get("conclusion") or {}
    H1("CONCLUSION & CALL TO ACTION")
    H2("1. Summary / Final Recap")
    paras(c.get("recap"))
    H2("2. Gospel Connection & Altar Call")
    if c.get("gospel_connection"):
        B(f"The Gospel Connection: {c['gospel_connection']}")
    if c.get("invitation"):
        B(f"The Invitation: {c['invitation']}")
    H2("3. Practical Next Steps (The Challenge)")
    if c.get("today"):
        B(c["today"] if c["today"].startswith("Immediate") else f"Immediate Response (Today): {c['today']}")
    if c.get("this_week"):
        B(c["this_week"] if c["this_week"].startswith("Weekly") else f"Weekly Integration (This Week): {c['this_week']}")
    rule()

    H1("SERMON LINES (TAKEAWAYS)")
    for t in o.get("takeaways") or []:
        B(t)
    rule()

    H1("KEY QUOTES & SOUNDBITES")
    for t in o.get("soundbites") or []:
        B(t)

    if o.get("preaching_notes"):
        rule()
        H1("NOTES FOR PREACHING")
        for t in o["preaching_notes"]:
            B(t)

    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(out))
    return out


def _missing(o: dict[str, Any]) -> list[str]:
    req = ["series", "title", "passage_ref", "passage_text", "goal", "intro", "points", "illustrations", "background",
           "conclusion", "takeaways", "soundbites"]
    miss = [k for k in req if not o.get(k)]
    intro = o.get("intro") or {}
    if intro and (not intro.get("hook") or len(intro.get("confusions") or []) < 3):
        miss.append("intro.hook / intro.confusions (need 3)")
    if len(o.get("points") or []) < 3:
        miss.append("points (need 3)")
    if len(o.get("illustrations") or []) < 2:
        miss.append("illustrations (need 2)")
    c = o.get("conclusion") or {}
    for k in ("recap", "gospel_connection", "invitation", "today", "this_week"):
        if c and not c.get(k):
            miss.append(f"conclusion.{k}")
    blob = json.dumps(o, ensure_ascii=False)
    if re.search(r"\[(FILL IN|TBD|INSERT|TODO)", blob, re.I):
        miss.append("placeholder text like [FILL IN] is not allowed — draft it and mark ✍️")
    return miss


@tool
def sermon_outline_schema() -> str:
    """Return the JSON schema (with an example) that compile_sermon_outline expects. Call this first, then fill EVERY
    field from the prep session before calling compile_sermon_outline."""
    return json.dumps(SCHEMA, ensure_ascii=False, indent=1)


@tool
def compile_sermon_outline(outline_json: str, upload: bool = True, account: str = "") -> str:
    """Render Pastor Stan's Sermon Outline Template as '<Title> - Sermon Outline.docx' (fixed layout, identical to the
    'It Ends at God' / 'Worth the Dig' documents) and upload it to Google Drive Church/Sermons. outline_json must follow
    sermon_outline_schema(): every section filled from the session in his own words; anything you draft yourself is
    prefixed with ✍️. Refuses to render if required sections are empty or contain [FILL IN]-style placeholders.
    Requires approval."""
    try:
        o = json.loads(outline_json)
    except Exception as e:  # noqa: BLE001
        return f"error: outline_json is not valid JSON ({e})"
    miss = _missing(o)
    if miss:
        return "error: outline incomplete — fill these from the session (draft + ✍️ if Stan never got to it): " + "; ".join(miss)
    name = f"{_slug(o['title'].replace('✍️', ''))} - Sermon Outline.docx"
    out = settings.workspace_dir / "sermons" / name
    try:
        render_docx(o, out)
    except Exception as e:  # noqa: BLE001
        return f"error: could not render docx: {type(e).__name__}: {e}"
    msg = f"rendered {out.relative_to(settings.workspace_dir)}"
    if not upload:
        return msg
    try:
        from app.tools import google_ws as gw

        email = gw._resolve(account)
        svc = gw._svc(email, "drive", "v3")
        folder_id = None
        parent = "root"
        for part in SERMONS_FOLDER_PATH.split("/"):
            q = (f"name = '{part}' and mimeType = 'application/vnd.google-apps.folder' and trashed = false"
                 + (f" and '{parent}' in parents" if parent != "root" else ""))
            res = svc.files().list(q=q, fields="files(id,name)", pageSize=5, supportsAllDrives=True,
                                   includeItemsFromAllDrives=True).execute().get("files", [])
            if not res:
                return msg + f" — but Drive folder '{SERMONS_FOLDER_PATH}' was not found ({email}); use drive_find_folder + drive_upload"
            parent = folder_id = res[0]["id"]
        from googleapiclient.http import MediaFileUpload

        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        f = svc.files().create(body={"name": name, "parents": [folder_id]}, media_body=MediaFileUpload(str(out), mimetype=mime),
                               fields="id,name,webViewLink", supportsAllDrives=True).execute()
        return msg + f"; uploaded to Drive {SERMONS_FOLDER_PATH} ({email}): {f.get('webViewLink')} (id={f['id']})"
    except Exception as e:  # noqa: BLE001
        return msg + f" — Drive upload failed: {type(e).__name__}: {e}"


sermon_outline_schema.metadata = {"requires_approval": False}
compile_sermon_outline.metadata = {"requires_approval": True}
TOOLS = [sermon_outline_schema, compile_sermon_outline]
